from __future__ import annotations

import logging
import re
import unicodedata
from pathlib import Path
from typing import Any, Iterator

from docx.document import Document as _Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx import Document

from app.knowledge.parsing.engine.base import BaseParser
from app.knowledge.parsing.enums import ParserFormat
from app.knowledge.parsing.exceptions import (
    FileParsingError,
    SourceFileNotFoundError,
)
from app.knowledge.parsing.schemas import NormalizedDocument, ParsingRequest

logger = logging.getLogger(__name__)


class DOCXParser(BaseParser):
    

    parser_name = "docx_parser"
    supported_extensions = (".docx",)

    MAX_FILE_SIZE_BYTES = 256 * 1024 * 1024
    MAX_BLOCKS = 100_000
    MAX_TABLE_CELLS = 1_000_000
    MAX_CELL_TEXT_CHARS = 10_000_000

    _MULTI_SPACE_RE = re.compile(r"[ \t]+")
    _MULTI_NEWLINE_RE = re.compile(r"\n{3,}")
    _SPACE_BEFORE_PUNCT_RE = re.compile(r"[ \t]+([,.;:!?%)\]}])")
    _SPACE_AFTER_OPEN_RE = re.compile(r"([(\[{])[ \t]+")
    _ZERO_WIDTH_RE = re.compile(r"[\u200b\u200c\u200d\ufeff]")

    def parse(
        self,
        file_path: Path,
        request: ParsingRequest,
    ) -> NormalizedDocument:
        self._validate_source(file_path)
        self._validate_request(request)

        try:
            document = Document(str(file_path))

            blocks = self._extract_blocks(document)
            content = self._join_content(
                block["content"] for block in blocks if block["content"]
            )

            if not content:
                raise FileParsingError(
                    f"DOCX contains no extractable text: {file_path.name}"
                )

         
            logical_page = {
                "page_number": 1,
                "content": content,
                "character_count": len(content),
                "word_count": self._count_words(content),
                "extraction_strategy": "docx_logical_document",
            }

            metadata = self._build_metadata(
                document=document,
                file_path=file_path,
                blocks=blocks,
                page=logical_page,
            )

            return NormalizedDocument(
                document_id=request.document_id,
                file_name=file_path.name,
                source_format=ParserFormat.DOCX,
                content=content,
                metadata=metadata,
                page_count=1,
                character_count=len(content),
                word_count=self._count_words(content),
                mime_type=request.mime_type,
            )

        except FileParsingError:
            raise
        except PackageNotFoundError as exc:
            raise FileParsingError(
                f"Invalid or corrupt DOCX package: {file_path.name}"
            ) from exc
        except Exception as exc:
            logger.exception(
                "DOCX parsing failed",
                extra={
                    "document_id": request.document_id,
                    "file_name": file_path.name,
                },
            )
            raise FileParsingError(
                f"Failed to parse DOCX: {file_path.name}"
            ) from exc

    # Validation

    def _validate_source(self, file_path: Path) -> None:
        if not file_path.exists():
            raise SourceFileNotFoundError(
                f"Source file not found: {file_path}"
            )

        if not file_path.is_file():
            raise SourceFileNotFoundError(
                f"Invalid source file: {file_path}"
            )

        if file_path.suffix.lower() != ".docx":
            raise FileParsingError(
                f"DOCX parser received non-DOCX file: {file_path.name}"
            )

        size = file_path.stat().st_size

        if size <= 0:
            raise FileParsingError(
                f"DOCX file is empty: {file_path.name}"
            )

        if size > self.MAX_FILE_SIZE_BYTES:
            raise FileParsingError(
                f"DOCX exceeds parser file-size limit "
                f"({self.MAX_FILE_SIZE_BYTES} bytes): {file_path.name}"
            )

    @staticmethod
    def _validate_request(request: ParsingRequest) -> None:
        if request.detected_format != ParserFormat.DOCX:
            raise FileParsingError(
                "DOCX parser received a request with "
                f"detected_format={request.detected_format!s}"
            )

        if request.extension.lower() != ".docx":
            raise FileParsingError(
                "DOCX parser received a request with "
                f"extension={request.extension!r}"
            )

        if not request.original_filename.strip():
            raise FileParsingError("Missing original filename")

    # Canonical block extraction

    @staticmethod
    def _iter_body_blocks(
        document: _Document,
    ) -> Iterator[Paragraph | Table]:
        body = document.element.body

        for child in body.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, document)
            elif isinstance(child, CT_Tbl):
                yield Table(child, document)

    def _extract_blocks(
        self,
        document: _Document,
    ) -> list[dict[str, Any]]:
        blocks: list[dict[str, Any]] = []

        for index, element in enumerate(
            self._iter_body_blocks(document),
            start=1,
        ):
            if index > self.MAX_BLOCKS:
                raise FileParsingError(
                    f"DOCX exceeds supported block limit "
                    f"({self.MAX_BLOCKS})"
                )

            if isinstance(element, Paragraph):
                block = self._paragraph_block(element, index)
                if block is not None:
                    blocks.append(block)
                continue

            block = self._table_block(element, index)
            blocks.append(block)

        return blocks

    def _paragraph_block(
        self,
        paragraph: Paragraph,
        index: int,
    ) -> dict[str, Any] | None:
        text = self._normalize_text(paragraph.text)

        if not text:
            return None

        style = self._safe_style_name(paragraph)
        heading_level = self._heading_level(style)
        list_info = self._list_info(paragraph)

        if heading_level is not None:
            block_type = "heading"
        elif list_info is not None:
            block_type = "list_item"
        else:
            block_type = "paragraph"

        block: dict[str, Any] = {
            "block_id": f"block-{index:06d}",
            "block_type": block_type,
            "content": text,
            "character_count": len(text),
            "word_count": self._count_words(text),
            "position": {"index": index - 1},
            "metadata": {},
        }

        if style:
            block["metadata"]["style"] = style

        if heading_level is not None:
            block["metadata"]["heading_level"] = heading_level

        if list_info is not None:
            block["metadata"]["list"] = list_info

        return block

    def _table_block(
        self,
        table: Table,
        index: int,
    ) -> dict[str, Any]:
        rows: list[list[str]] = []
        cell_count = 0

        for row in table.rows:
            values: list[str] = []

            for cell in row.cells:
                cell_count += 1

                if cell_count > self.MAX_TABLE_CELLS:
                    raise FileParsingError(
                        f"DOCX exceeds supported table-cell limit "
                        f"({self.MAX_TABLE_CELLS})"
                    )

                text = self._extract_cell_text(cell)

                if len(text) > self.MAX_CELL_TEXT_CHARS:
                    raise FileParsingError(
                        "DOCX table cell exceeds supported text limit"
                    )

                values.append(text)

            rows.append(values)

        table_text = self._table_to_text(rows)

        return {
            "block_id": f"block-{index:06d}",
            "block_type": "table",
            "content": table_text,
            "character_count": len(table_text),
            "word_count": self._count_words(table_text),
            "position": {"index": index - 1},
            "metadata": {
                "rows": rows,
                "row_count": len(rows),
                "column_count": max(
                    (len(row) for row in rows),
                    default=0,
                ),
            },
        }

    @classmethod
    def _extract_cell_text(cls, cell: _Cell) -> str:
        parts: list[str] = []

        for paragraph in cell.paragraphs:
            text = cls._normalize_text(paragraph.text)
            if text:
                parts.append(text)

        return "\n".join(parts)

    @staticmethod
    def _table_to_text(rows: list[list[str]]) -> str:
        return "\n".join(
            " | ".join(cell.replace("\n", " ").strip() for cell in row)
            for row in rows
            if any(cell.strip() for cell in row)
        ).strip()

    # Text normalization

    @classmethod
    def _normalize_text(cls, text: str) -> str:
        if not text:
            return ""

        text = unicodedata.normalize("NFKC", text)
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        text = text.replace("\u00a0", " ")
        text = cls._ZERO_WIDTH_RE.sub("", text)

        lines: list[str] = []

        for raw_line in text.split("\n"):
            line = raw_line.strip()

            if not line:
                continue

            line = cls._MULTI_SPACE_RE.sub(" ", line)
            line = cls._SPACE_BEFORE_PUNCT_RE.sub(r"\1", line)
            line = cls._SPACE_AFTER_OPEN_RE.sub(r"\1", line)

            lines.append(line)

        return cls._MULTI_NEWLINE_RE.sub(
            "\n\n",
            "\n".join(lines),
        ).strip()

    @staticmethod
    def _join_content(parts: Any) -> str:
        cleaned = [part.strip() for part in parts if part and part.strip()]
        return "\n\n".join(cleaned)

    @staticmethod
    def _count_words(text: str) -> int:
        return len(re.findall(r"\S+", text))

    # Paragraph metadata

    @staticmethod
    def _safe_style_name(paragraph: Paragraph) -> str | None:
        try:
            value = paragraph.style.name
        except Exception:
            return None

        return str(value) if value else None

    @staticmethod
    def _heading_level(style_name: str | None) -> int | None:
        if not style_name:
            return None

        match = re.fullmatch(
            r"Heading\s+([1-9][0-9]*)",
            style_name.strip(),
            flags=re.IGNORECASE,
        )

        return int(match.group(1)) if match else None

    @staticmethod
    def _list_info(
        paragraph: Paragraph,
    ) -> dict[str, Any] | None:
        p_pr = paragraph._p.pPr

        if p_pr is None or p_pr.numPr is None:
            return None

        num_pr = p_pr.numPr

        if num_pr.numId is None:
            return None

        return {
            "num_id": int(num_pr.numId.val),
            "level": (
                int(num_pr.ilvl.val)
                if num_pr.ilvl is not None
                else 0
            ),
        }

    # Metadata

    def _build_metadata(
        self,
        document: _Document,
        file_path: Path,
        blocks: list[dict[str, Any]],
        page: dict[str, Any],
    ) -> dict[str, Any]:
        core = document.core_properties

        core_properties = {
            "title": core.title or None,
            "subject": core.subject or None,
            "author": core.author or None,
            "keywords": core.keywords or None,
            "comments": core.comments or None,
            "category": core.category or None,
            "created": (
                core.created.isoformat()
                if core.created
                else None
            ),
            "modified": (
                core.modified.isoformat()
                if core.modified
                else None
            ),
            "last_modified_by": (
                core.last_modified_by or None
            ),
        }

        try:
            inline_shape_count = len(document.inline_shapes)
        except Exception:
            inline_shape_count = 0

        return {
            "file_name": file_path.name,
            "file_extension": file_path.suffix.lower(),
            "parser": self.parser_name,
            "format": ParserFormat.DOCX.value,
            "file_size_bytes": file_path.stat().st_size,

            "page_count": 1,
            "pages": [page],

            "blocks": blocks,

            "paragraph_count": sum(
                block["block_type"] == "paragraph"
                for block in blocks
            ),
            "heading_count": sum(
                block["block_type"] == "heading"
                for block in blocks
            ),
            "list_count": sum(
                block["block_type"] == "list_item"
                for block in blocks
            ),
            "table_count": sum(
                block["block_type"] == "table"
                for block in blocks
            ),
            "inline_shape_count": inline_shape_count,
            "core_properties": core_properties,
            "extraction_strategies": [
                "docx_logical_document"
            ],
            "warnings": [],
        }
